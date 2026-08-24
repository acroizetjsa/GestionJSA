#!/usr/bin/env python3
"""Export personnel des parcs bus/cars publics de TC Infos vers Excel.

Le script consulte uniquement les pages publiques, temporise les requetes et
ne contourne aucune authentification ou protection. Il essaie d'abord le site
`dev.tc-infos.fr`, qui est actuellement joignable depuis les moteurs/robots,
puis le domaine principal en secours.
"""
from __future__ import annotations

import argparse
import re
import time
from datetime import datetime, timezone
from pathlib import Path
from typing import Iterable

import requests
from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BASE_URLS = ["https://dev.tc-infos.fr", "https://tc-infos.fr"]
USER_AGENT = "Mozilla/5.0 (compatible; PersonalDataExport/1.1; +https://github.com/acroizetjsa/GestionJSA)"
PAGE_SIZE = 200
MAX_PAGES_PER_NETWORK = 30

HEADERS = [
    "ReseauId", "Reseau", "NumeroParc", "Constructeur", "Modele",
    "Immatriculation", "MiseEnCirculation", "ArriveeReseau",
    "Exploitant", "Depot", "SourceURL", "DateExtractionUTC",
]


def clean(text: str) -> str:
    return re.sub(r"\s+", " ", text.replace("\xa0", " ")).strip()


def network_name(soup: BeautifulSoup, network_id: int) -> str:
    h1 = soup.find("h1") or soup.find("h2")
    title = clean(h1.get_text(" ", strip=True)) if h1 else ""
    title = re.sub(r"^Parc bus\s*/\s*cars\s+", "", title, flags=re.I)
    if title:
        return title
    page_title = clean(soup.title.get_text(" ", strip=True)) if soup.title else ""
    page_title = re.sub(r"^Parc bus\s*/\s*cars\s+", "", page_title, flags=re.I)
    page_title = re.sub(r"\s*-\s*TC Infos\s*$", "", page_title, flags=re.I)
    return page_title or f"Reseau {network_id}"


def parse_table(soup: BeautifulSoup, network_id: int, source_url: str) -> list[dict[str, str]]:
    table = soup.find("table")
    if table is None:
        return []
    rows = table.find_all("tr")
    if not rows:
        return []

    header_cells = rows[0].find_all(["th", "td"])
    headers = [clean(c.get_text(" ", strip=True)).lower() for c in header_cells]
    aliases = {
        "numero": ["n°", "nº", "numero", "numéro"],
        "constructeur": ["constructeur"],
        "modele": ["modèle", "modele"],
        "immat": ["immatriculation"],
        "mec": ["mise en circulation"],
        "arrivee": ["arrivée sur le réseau", "arrivee sur le reseau", "arrivée réseau"],
        "exploitant": ["exploitant"],
        "depot": ["dépôt", "depot"],
    }

    def idx_for(keys: Iterable[str]) -> int | None:
        for i, h in enumerate(headers):
            if h in keys:
                return i
        return None

    idx = {k: idx_for(v) for k, v in aliases.items()}
    if idx["constructeur"] is None or idx["modele"] is None:
        # Ordre public TC Infos lorsque le balisage d'entete n'est pas standard.
        idx = {"numero": 0, "constructeur": 1, "modele": 2, "immat": 3,
               "mec": 4, "arrivee": 5, "exploitant": 6, "depot": 7}
        data_rows = rows
    else:
        data_rows = rows[1:]

    name = network_name(soup, network_id)
    extracted = datetime.now(timezone.utc).replace(microsecond=0).isoformat()
    result: list[dict[str, str]] = []

    for tr in data_rows:
        cells = tr.find_all("td")
        if not cells:
            continue
        values = [clean(c.get_text(" ", strip=True)) for c in cells]

        def val(key: str) -> str:
            i = idx.get(key)
            return values[i] if i is not None and i < len(values) else ""

        constructeur, modele = val("constructeur"), val("modele")
        immat, numero = val("immat"), val("numero")
        if not any([constructeur, modele, immat, numero]):
            continue
        if constructeur.lower() == "constructeur" or modele.lower() in {"modèle", "modele"}:
            continue

        result.append({
            "ReseauId": str(network_id),
            "Reseau": name,
            "NumeroParc": numero,
            "Constructeur": constructeur,
            "Modele": modele,
            "Immatriculation": immat,
            "MiseEnCirculation": val("mec"),
            "ArriveeReseau": val("arrivee"),
            "Exploitant": val("exploitant"),
            "Depot": val("depot"),
            "SourceURL": source_url,
            "DateExtractionUTC": extracted,
        })
    return result


def request_page(session: requests.Session, network_id: int, page: int, timeout: float) -> tuple[int, list[dict[str, str]], str]:
    suffix = "" if page == 1 else f"/page{page}"
    errors: list[str] = []
    saw_404 = False

    for base in BASE_URLS:
        url = f"{base}/reseau/{network_id}/parc-bus-cars{suffix}"
        for attempt in range(2):
            try:
                r = session.get(url, timeout=timeout, allow_redirects=True)
                if r.status_code == 404:
                    saw_404 = True
                    break
                if r.status_code == 429:
                    wait = 4 * (attempt + 1)
                    print(f"[{network_id} p{page}] 429 sur {base}; pause {wait}s")
                    time.sleep(wait)
                    continue
                r.raise_for_status()
                soup = BeautifulSoup(r.text, "html.parser")
                return r.status_code, parse_table(soup, network_id, url), ""
            except requests.RequestException as exc:
                errors.append(f"{base}: {type(exc).__name__}: {exc}")
                if attempt == 0:
                    time.sleep(1.5)
        # Essayer l'hote suivant.

    if saw_404 and not errors:
        return 404, [], ""
    return 0, [], " | ".join(errors[-4:])


def fetch_network(session: requests.Session, network_id: int, timeout: float) -> tuple[int, list[dict[str, str]], str]:
    all_rows: list[dict[str, str]] = []
    for page in range(1, MAX_PAGES_PER_NETWORK + 1):
        status, rows, error = request_page(session, network_id, page, timeout)
        if status == 0:
            return 0, all_rows, error
        if status == 404:
            return (404 if page == 1 else 200), all_rows, ""
        if not rows:
            return 200, all_rows, ""
        all_rows.extend(rows)
        # Les listes TC Infos sont paginees par 200 vehicules.
        if len(rows) < PAGE_SIZE:
            return 200, all_rows, ""
    return 200, all_rows, "limite de pagination atteinte"


def autosize(ws) -> None:
    widths: dict[int, int] = {}
    for row in ws.iter_rows(values_only=True):
        for i, value in enumerate(row, 1):
            if value is not None:
                widths[i] = min(max(widths.get(i, 0), len(str(value)) + 2), 60)
    for i, width in widths.items():
        ws.column_dimensions[get_column_letter(i)].width = max(width, 12)


def write_xlsx(rows: list[dict[str, str]], output: Path, stats: dict[str, int]) -> int:
    wb = Workbook()
    ws = wb.active
    ws.title = "Vehicules"
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}1"
    ws.sheet_view.showGridLines = False
    header_fill = PatternFill("solid", fgColor="1F4E78")

    for col, title in enumerate(HEADERS, 1):
        c = ws.cell(1, col, title)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = header_fill
        c.alignment = Alignment(horizontal="center", vertical="center")

    seen = set()
    deduped: list[dict[str, str]] = []
    for row in rows:
        # Immatriculation est le meilleur identifiant public; conserver le reseau
        # dans la cle pour ne pas perdre un cas ambigu ou une immatriculation vide.
        key = (row["ReseauId"], row["Immatriculation"], row["NumeroParc"], row["Constructeur"], row["Modele"])
        if key in seen:
            continue
        seen.add(key)
        deduped.append(row)

    for r_idx, row in enumerate(deduped, 2):
        for c_idx, field in enumerate(HEADERS, 1):
            cell = ws.cell(r_idx, c_idx, row.get(field, ""))
            cell.font = Font(color="008000")
            cell.alignment = Alignment(vertical="top")
    autosize(ws)

    meta = wb.create_sheet("Extraction")
    meta.sheet_view.showGridLines = False
    metadata = [
        ("Sources", " ; ".join(BASE_URLS)),
        ("Usage", "Export personnel de pages publiques"),
        ("Date extraction UTC", datetime.now(timezone.utc).replace(microsecond=0).isoformat()),
        ("Reseaux testes", stats.get("tested", 0)),
        ("Reseaux avec vehicules", stats.get("with_rows", 0)),
        ("Erreurs transport", stats.get("transport_errors", 0)),
        ("Lignes collectees avant dedoublonnage", len(rows)),
        ("Vehicules/lignes apres dedoublonnage", len(deduped)),
        ("Regle", "Aucune authentification/protection contournee; temporisation entre requetes."),
    ]
    for i, (k, v) in enumerate(metadata, 1):
        meta.cell(i, 1, k).font = Font(bold=True, color="666666")
        meta.cell(i, 2, v).font = Font(color="666666")
    autosize(meta)

    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)
    return len(deduped)


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--start-network-id", type=int, default=1)
    ap.add_argument("--max-network-id", type=int, default=2200)
    ap.add_argument("--delay", type=float, default=0.5)
    ap.add_argument("--timeout", type=float, default=15.0)
    ap.add_argument("--output", default="tc_infos_vehicules.xlsx")
    ap.add_argument("--fail-if-empty", action="store_true")
    args = ap.parse_args()

    session = requests.Session()
    session.headers.update({
        "User-Agent": USER_AGENT,
        "Accept-Language": "fr-FR,fr;q=0.9,en;q=0.6",
        "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8",
        "Connection": "keep-alive",
    })

    rows: list[dict[str, str]] = []
    stats = {"tested": 0, "with_rows": 0, "transport_errors": 0}
    consecutive_transport_errors = 0

    for network_id in range(args.start_network_id, args.max_network_id + 1):
        stats["tested"] += 1
        status, found, error = fetch_network(session, network_id, args.timeout)
        if status == 0:
            stats["transport_errors"] += 1
            consecutive_transport_errors += 1
            print(f"[{network_id}] erreur transport: {error}")
            if consecutive_transport_errors >= 10:
                print("10 erreurs transport consecutives: arret de securite.")
                break
        else:
            consecutive_transport_errors = 0

        if found:
            rows.extend(found)
            stats["with_rows"] += 1
            print(f"[{network_id}] {len(found)} lignes; total={len(rows)}")
        elif network_id % 50 == 0:
            print(f"[{network_id}] progression; total={len(rows)}")
        time.sleep(max(args.delay, 0.25))

    count = write_xlsx(rows, Path(args.output), stats)
    print(f"Export termine: {args.output} ({len(rows)} lignes brutes, {count} apres dedoublonnage)")
    if args.fail_if_empty and count == 0:
        print("ERREUR: aucune ligne extraite.")
        return 2
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
